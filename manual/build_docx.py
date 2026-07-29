"""Genera el manual de usuario en formato Word (.docx) con las capturas reales."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMGS = "/app/manual/imgs"
OUT = "/app/manual/manual-usuario-megasoft.docx"

doc = Document()

# --- estilos base ---
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

sect = doc.sections[0]
sect.top_margin = Cm(2)
sect.bottom_margin = Cm(2)
sect.left_margin = Cm(2.2)
sect.right_margin = Cm(2.2)


def shade_paragraph(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def h1(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)


def h2(num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"  {num}  ")
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0xFA, 0xCC, 0x15)
    r1.font.size = Pt(14)
    shade_paragraph(p, "0F172A")  # highlight the number bg via paragraph shade — imperfect but readable
    r2 = p.add_run(f"  {text}")
    r2.font.bold = True
    r2.font.size = Pt(16)
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def h3(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)


def body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(it)


def steps(items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(it)


def add_image(path, caption, width_in=6.0):
    doc.add_picture(path, width=Inches(width_in))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def callout(text, kind="tip"):
    palette = {
        "tip":  ("ECFDF5", 0x06, 0x4E, 0x3B),
        "warn": ("FEE2E2", 0x7F, 0x1D, 0x1D),
        "note": ("FEF9C3", 0x71, 0x3F, 0x12),
    }
    fill, cr, cg, cb = palette[kind]
    p = doc.add_paragraph()
    shade_paragraph(p, fill)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(cr, cg, cb)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


# ================================================================
# COVER
# ================================================================
sub = doc.add_paragraph()
sr = sub.add_run("MEGASOFT · SUITE CORPORATIVA DE ASISTENCIA")
sr.font.size = Pt(10)
sr.font.bold = True
sr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

h1("Manual rápido de usuario")

kk = doc.add_paragraph()
kr = kk.add_run("Cómo iniciar sesión, cambiar tu contraseña y registrar tu rostro · v1.0")
kr.font.size = Pt(11)
kr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

# separator
sep = doc.add_paragraph()
sep_r = sep.add_run("_" * 90)
sep_r.font.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)

# ================================================================
# 1. LOGIN
# ================================================================
h2("1", "Cómo entrar al sistema")
body("Abre desde cualquier navegador la dirección corporativa: "
     "https://asistencia-web-1.emergent.host. Verás la pantalla de acceso corporativo.")

add_image(f"{IMGS}/01-login.jpeg",
          "Pantalla principal — panel izquierdo con la marca MegaSoft, formulario de acceso a la derecha.")

h3("Pasos")
steps([
    "Escribe tu correo electrónico corporativo (ej. tunombre@megasoft.com.ve).",
    "Escribe tu contraseña. Puedes tocar el ícono del ojo para mostrarla si necesitas revisarla.",
    "Presiona el botón 'Entrar al panel'.",
])

add_image(f"{IMGS}/02-login-fill.jpeg",
          "Con el correo y contraseña completos, presiona 'Entrar al panel'.")

callout("Si olvidas la contraseña o el sistema no te reconoce, contacta al administrador de RRHH o "
        "de sistemas — ellos pueden restablecerla desde la sección de Empleados.", "note")

doc.add_page_break()

# ================================================================
# 2. CHANGE PASSWORD
# ================================================================
h2("2", "Cómo cambiar tu contraseña")
body("Una vez adentro, tu nombre aparece en la esquina superior derecha. Haz clic sobre tu avatar "
     "(círculo con tus iniciales) para abrir el menú de usuario.")

add_image(f"{IMGS}/03-usermenu.jpeg",
          "Menú superior derecho — muestra tu correo en sesión y las opciones disponibles.")

h3("Pasos")
steps([
    "Haz clic en tu avatar / nombre arriba a la derecha.",
    "Selecciona 'Cambiar contraseña'.",
    "Completa el formulario: contraseña actual, nueva contraseña (mínimo 8 caracteres) y confirmación.",
    "Presiona 'Actualizar'. Verás una confirmación y podrás seguir usando el sistema.",
])

h3("Requisitos de la nueva contraseña")
bullets([
    "Mínimo 8 caracteres.",
    "Distinta a la contraseña actual.",
    "Se recomienda combinar mayúsculas, números y símbolos.",
])

add_image(f"{IMGS}/04-changepw.jpeg",
          "Diálogo 'Cambiar contraseña' — valida los 3 campos antes de habilitar el botón.")

callout("Consejo: cambia tu contraseña al menos cada 90 días y nunca la compartas. "
        "Ninguna persona del área técnica te la va a pedir.", "tip")

doc.add_page_break()

# ================================================================
# 3. REGISTER FACE
# ================================================================
h2("3", "Cómo registrar tu rostro")
body("El sistema usa reconocimiento facial en el Kiosco de asistencia para registrar tu entrada y salida "
     "en un segundo — sin apps, sin contraseñas. Necesitas registrar tu rostro una sola vez.")

h3("Cómo llegar")
steps([
    "Ingresa al sistema con tu usuario y contraseña.",
    "En el menú superior derecho (tu avatar), haz clic en 'Registrar rostro'. También puedes acceder directo por /onboarding.",
    "El navegador te pedirá permiso para usar la cámara — acepta.",
])

add_image(f"{IMGS}/05-onboarding.jpeg",
          "Pantalla 'Registra tu rostro' — verás tu cámara activa al lado derecho.")

h3("Recomendaciones antes de la captura")
bullets([
    "Luz frontal y natural · evita luces detrás de ti.",
    "Solo tu rostro en cuadro · sin gorra, lentes oscuros ni mascarilla.",
    "Mira directo a la cámara · expresión neutral, sin gestos exagerados.",
    "Encuadra tu cara dentro del recuadro con líneas punteadas.",
])

h3("Pasos")
steps([
    "Alinea tu rostro dentro del recuadro verde.",
    "Presiona el botón 'Capturar'.",
    "Si la calidad de la foto es baja o no se detecta rostro, el sistema te pedirá reintentar.",
    "Cuando reciba la confirmación 'Rostro registrado', ya podrás marcar entrada y salida desde el kiosco compartido de tu sede acercándote a la cámara.",
])

callout("Importante: la selfie se guarda cifrada en el servidor y nunca se comparte. "
        "Solo se usa para verificar tu identidad al marcar en el kiosco.", "warn")

callout("Si cambias tu apariencia (barba, corte de pelo, lentes nuevos permanentes), "
        "simplemente vuelve a 'Registrar rostro' desde tu menú de usuario — se actualiza tu plantilla biométrica.", "tip")

# footer
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = foot.add_run("© MegaSoft Computación, C.A. — Todos los derechos reservados. Documento generado para uso interno.")
fr.font.size = Pt(9)
fr.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

doc.save(OUT)
print(f"OK → {OUT}")
