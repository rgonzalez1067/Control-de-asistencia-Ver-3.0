"""Genera los manuales Word (Supervisor + Asignación de Horarios).

Uso:
    python3 /app/manual/build_manuals_v2.py
Produce:
    /app/manual/Manual_Supervisor_MegaSoft_Asistencia.docx
    /app/manual/Manual_Asignacion_Horarios_MegaSoft.docx
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SUP = "/app/manual/screens_supervisor"
ASG = "/app/manual/screens_asignacion"

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
AMBER = RGBColor(0xC7, 0x8A, 0x1F)
RED = RGBColor(0xB0, 0x2A, 0x2A)
EMERALD = RGBColor(0x1F, 0x7A, 0x4C)
GREY = RGBColor(0x55, 0x5D, 0x6A)


def _shade(cell, hex_color):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


def r(p, t, bold=False, size=None, color=None, italic=False):
    run = p.add_run(t)
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.font.name = "Calibri"
    return run


def h(doc, t, level=1, color=NAVY):
    hh = doc.add_heading("", level=level)
    r(hh, t, color=color)
    return hh


def p(doc, t=None, bold=False, size=11, color=None, italic=False, align=None):
    pp = doc.add_paragraph()
    if align is not None:
        pp.alignment = align
    if t:
        r(pp, t, bold=bold, size=size, color=color, italic=italic)
    return pp


def b(doc, t, level=0):
    pp = doc.add_paragraph(style="List Bullet")
    pp.paragraph_format.left_indent = Cm(0.5 + level * 0.6)
    r(pp, t, size=11)


def n(doc, t):
    pp = doc.add_paragraph(style="List Number")
    r(pp, t, size=11)


def img(doc, path, cm=15.5, cap=None):
    pp = doc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(path, width=Cm(cm))
    if cap:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r(c, f"Figura: {cap}", size=9, italic=True, color=GREY)


def callout(doc, title, lines, color_hex="FFF6E5", accent=AMBER):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    _shade(c, color_hex)
    p1 = c.paragraphs[0]
    r(p1, "► " + title, bold=True, size=11, color=accent)
    for line in lines:
        pp = c.add_paragraph()
        r(pp, line, size=10.5)


def pb(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def cover(doc, main_title, subtitle, badge):
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r(c, "MegaSoft Computación, C.A.\n", bold=True, size=14, color=NAVY)
    r(c, "Suite Corporativa de Asistencia\n\n\n", size=12, color=GREY)
    r(c, "\n\n" + main_title + "\n\n", bold=True, size=26, color=NAVY)
    r(c, subtitle + "\n", size=14, color=AMBER, italic=True)
    r(c, badge + "\n", size=12, color=GREY)
    r(c, "\n\n\n\n\n\n\n\n\n\nVersión 1.0 — Febrero 2026\n", size=11, color=GREY)


def setup(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)


# =============================================================
# MANUAL 1 — SUPERVISOR
# =============================================================
sup = Document()
setup(sup)
cover(sup,
      "MANUAL DEL SUPERVISOR",
      "Consulta, análisis y aprobación de tu equipo",
      "Dirigido a: Coordinadores, Gerentes y Directores")
pb(sup)

h(sup, "Contenido", 1)
for t in [
    "1.  Introducción y perfil de usuario",
    "2.  Cómo iniciar sesión",
    "3.  Mi Equipo — Panel operativo",
    "     3.1  KPIs y matriz diaria",
    "     3.2  Aprobar o rechazar una justificación (paso a paso)",
    "     3.3  Asignar el horario base de un miembro",
    "4.  Matriz de Asistencia — Análisis avanzado",
    "     4.1  Filtros y tipos de horario",
    "     4.2  Interpretación de columnas y colores",
    "     4.3  Exportar PDF / Excel",
    "5.  Menú Reportes — Auditoría y descargas",
    "     5.1  Filtros disponibles",
    "     5.2  Exportar CSV para nómina o RRHH",
    "6.  Novedades del equipo",
    "     6.1  Pestañas Pendientes · Historial · Todas",
    "     6.2  Aprobar y rechazar en bloque",
    "     6.3  Registrar una novedad a nombre de un colaborador",
    "7.  Buenas prácticas del supervisor",
    "8.  Preguntas frecuentes",
]:
    r(sup.add_paragraph(), t, size=11)
pb(sup)

h(sup, "1. Introducción y perfil de usuario", 1)
p(sup, "Este manual está dirigido al personal con rol de liderazgo (Coordinador, Gerente y Director). "
       "Un supervisor tiene acceso a las mismas pantallas que un empleado y, adicionalmente, a las "
       "vistas de gestión y aprobación de su equipo directo.", size=11)
p(sup, "Como supervisor podrás:", size=11, bold=True)
b(sup, "Consultar la asistencia de tu equipo día a día (Mi Equipo).")
b(sup, "Aprobar o rechazar las justificaciones de tardanzas que envíen tus colaboradores.")
b(sup, "Analizar la Matriz de Asistencia consolidada para tu equipo o para toda la organización.")
b(sup, "Descargar reportes en CSV / PDF / Excel para nómina y RRHH.")
b(sup, "Aprobar novedades (vacaciones, reposos, permisos) y registrar novedades a nombre de tus colaboradores.")
b(sup, "Asignar el horario base a cada miembro de tu equipo directamente desde Mi Equipo.")
callout(sup, "REGLA DE ALCANCE (SCOPE)",
        ["Un supervisor sólo puede ver, aprobar o rechazar registros de empleados que lo tengan asignado como supervisor directo (campo supervisor_id).",
         "Un rol Administrador tiene acceso global sin límite de scope.",
         "El acceso a cada opción depende de tu Perfil de Acceso (RBAC). Si no ves alguna sección aquí descrita, solicita permiso a tu administrador."],
        "EEF5FF", NAVY)
pb(sup)

h(sup, "2. Cómo iniciar sesión", 1)
n(sup, "Abre la URL corporativa en tu navegador.")
n(sup, "Ingresa tu correo institucional y contraseña.")
n(sup, "Presiona 'Entrar al panel'.")
p(sup, "Si tu usuario tiene múltiples permisos (por ejemplo, además de coordinador eres administrador), verás todas las opciones habilitadas en el menú lateral.", size=11)
pb(sup)

# ------- 3. Mi Equipo -------
h(sup, "3. Mi Equipo — Panel operativo", 1)
p(sup, "Es la pantalla del día a día del supervisor. Muestra en una única matriz a todos tus colaboradores con sus marcajes en los últimos N días.", size=11)
img(sup, f"{SUP}/01_team_top.jpg", 16.5, "Mi Equipo: KPIs superiores + matriz diaria de marcajes.")

h(sup, "3.1 KPIs y matriz diaria", 2)
b(sup, "Miembros: total de personas bajo tu supervisión (excluye admins si eres admin viendo la vista global).")
b(sup, "Entradas hoy: marcas de tipo Entrada en el día actual.")
b(sup, "Tarde en el período: cantidad de entradas late en el rango consultado.")
b(sup, "Justif. pendientes: registros con justificación enviada pero aún no decidida por ti.")
p(sup, "En la matriz, cada celda representa un día para un empleado:", size=11)
b(sup, "Verde: día normal, marcaje a tiempo.")
b(sup, "Ámbar: retraso leve.")
b(sup, "Rojo suave: retraso mayor.")
b(sup, "Ámbar pulsante con ícono de expediente: justificación pendiente de aprobación (clickeable).")
b(sup, "Verde intenso con check: justificación aprobada.")
b(sup, "Rojo con X: justificación rechazada.")

h(sup, "3.2 Aprobar o rechazar una justificación (paso a paso)", 2)
p(sup, "Esta es la operación más importante del supervisor. Cuando un colaborador envía una justificación por una tardanza, aparece un indicador visual en la celda del día justificado.", size=11)

img(sup, f"{SUP}/02_team_cell.jpg", 15.5, "Celda con justificación pendiente resaltada en la matriz.")

n(sup, "Ubica al colaborador en la lista y busca la celda del día justificado (aparece con fondo ámbar pulsante y un ícono de expediente ⚠).")
n(sup, "Haz clic sobre la celda. Se abrirá el diálogo 'Revisar justificación' con los datos de la solicitud.")
img(sup, f"{SUP}/03_team_dialog.jpg", 14.0, "Diálogo con la justificación del empleado.")

p(sup, "El diálogo muestra tres bloques de información:", size=11)
b(sup, "Minutos de retraso: cuánto se demoró el colaborador respecto a su hora esperada o su tiempo de descanso.")
b(sup, "Estado actual: badge que confirma que está Pendiente.")
b(sup, "Justificación del empleado: el texto que el colaborador ingresó en su Mi Historial.")

p(sup, "Tienes dos acciones posibles:", bold=True, size=11)

# Table for actions
t = sup.add_table(rows=3, cols=3)
t.style = "Light Grid Accent 1"
hd = t.rows[0].cells
r(hd[0].paragraphs[0], "Acción", bold=True, size=10, color=NAVY)
r(hd[1].paragraphs[0], "Efecto sobre el registro", bold=True, size=10, color=NAVY)
r(hd[2].paragraphs[0], "Efecto sobre reportes", bold=True, size=10, color=NAVY)

r1 = t.rows[1].cells
r(r1[0].paragraphs[0], "Aceptar justificación", size=10, bold=True, color=EMERALD)
r(r1[1].paragraphs[0], "Status → Retraso Justificado. La celda cambia a verde.", size=10)
r(r1[2].paragraphs[0], "Los minutos de retraso NO se contabilizan en 'Minutos Perdidos'.", size=10)

r2 = t.rows[2].cells
r(r2[0].paragraphs[0], "Rechazar", size=10, bold=True, color=RED)
r(r2[1].paragraphs[0], "Debes escribir una 'Razón del rechazo' obligatoria (mínimo 3 caracteres). Status → Retraso Injustificado. La celda cambia a rojo.", size=10)
r(r2[2].paragraphs[0], "Los minutos SÍ se suman a los 'Minutos Perdidos' del empleado.", size=10)

img(sup, f"{SUP}/04_team_reject.jpg", 14.0, "Flujo de rechazo con la razón obligatoria.")

callout(sup, "CÓMO REDACTAR UNA BUENA RAZÓN DE RECHAZO",
        ["Sé concreto y profesional. Ej.: 'La justificación no contiene evidencia y es la tercera ocasión en el mes con el mismo motivo. Se envía notificación al colaborador para reunión formal.'",
         "Evita frases como 'no acepto' sin explicación: el empleado leerá tu comentario en su Historial y necesita entender el criterio.",
         "Si tienes dudas, guarda la decisión pendiente y conversa primero con el colaborador. Una vez decidida (aprobada o rechazada), no se puede editar en el sistema."],
        "FDECEC", RED)

h(sup, "3.3 Asignar el horario base de un miembro", 2)
p(sup, "En la columna izquierda, cada miembro tiene un selector 'Día Completo / Turno Uno / …' que corresponde a su horario base. Elige el horario apropiado desde el desplegable; el cambio se guarda automáticamente.", size=11)
p(sup, "Nota: los cambios de horario base afectan sólo los cálculos futuros. Los marcajes ya registrados conservan la clasificación calculada en su momento.", size=10, italic=True, color=GREY)
pb(sup)

# ------- 4. Matriz de Asistencia -------
h(sup, "4. Matriz de Asistencia — Análisis avanzado", 1)
p(sup, "La Matriz de Asistencia (menú lateral 'Matriz de asistencia') es la vista consolidada para análisis, exportación y auditoría. Puedes filtrar por rango de fechas, tipo de horario, sede, departamento y empleado.", size=11)
img(sup, f"{SUP}/05_matricial_top.jpg", 16.5, "Reporte matricial de asistencia — barra de filtros.")

h(sup, "4.1 Filtros y tipos de horario", 2)
b(sup, "Desde / Hasta: rango de días a analizar (recomendado no exceder 62 días para conservar la velocidad de carga).")
b(sup, "Tipo de horario: filtra a los empleados según cuántos bloques laboran (1 bloque, 2 bloques) o si trabajan con Horario Especial (turnos rotativos).")
b(sup, "Sede: útil para supervisores multi-sede.")
b(sup, "Departamento / Empleado: acota aún más el resultado.")

h(sup, "4.2 Interpretación de columnas y colores", 2)
p(sup, "Cada empleado aparece en una fila y cada día en columna. Un día puede mostrar hasta 4 sub-columnas: E1 (primera entrada), S1 (primera salida), E2 (segunda entrada), S2 (segunda salida). Los colores indican:", size=11)
b(sup, "Hora en negro: dentro de la tolerancia. Todo bien.")
b(sup, "Hora en rojo: tardanza fuera de tolerancia (E1) o exceso de descanso >60 min (E2).")
b(sup, "Ámbar en cursiva: salida auto-imputada a las 23:59 (el sistema cerró el día porque no marcó salida).")
b(sup, "Azul suave 'VACACIONES / REPOSO / PERMISO': novedad de día completo aprobada.")
b(sup, "Rojo intenso 'FALTA': día laboral sin ningún marcaje ni novedad.")
b(sup, "Sub-fila 'Cita médica / Permiso / Visita': novedad parcial dentro de un día laborable.")
img(sup, f"{SUP}/06_matricial_body.jpg", 16.5, "Detalle de filas con marcajes, faltas y novedades.")

p(sup, "Columnas de totales al extremo derecho:", size=11)
b(sup, "MIN. PERD.: minutos perdidos acumulados por retrasos NO justificados.")
b(sup, "T-J: tardanzas justificadas (aceptadas por ti).")
b(sup, "T-NJ: tardanzas no justificadas o rechazadas.")
b(sup, "NOVEDADES: cantidad de novedades aprobadas en el período.")
b(sup, "FALTAS: días laborales sin marcaje ni novedad.")

h(sup, "4.3 Exportar PDF / Excel", 2)
p(sup, "Los botones 'PDF' y 'Excel' en la esquina superior derecha generan el archivo con los filtros aplicados. Úsalos para presentar a Recursos Humanos o adjuntar a evaluaciones de desempeño.", size=11)
pb(sup)

# ------- 5. Reportes -------
h(sup, "5. Menú Reportes — Auditoría y descargas", 1)
p(sup, "La sección 'Reportes' (menú lateral) muestra todos los marcajes crudos en una tabla ordenable, con filtros de rango, tipo, empleado, departamento y sede. Es ideal para auditar un caso puntual o exportar un CSV para nómina.", size=11)
img(sup, f"{SUP}/07_reportes_top.jpg", 16.5, "Menú Reportes — filtros y KPIs.")

h(sup, "5.1 Filtros disponibles", 2)
b(sup, "Fecha desde/hasta.")
b(sup, "Tipo: Entrada / Salida.")
b(sup, "Empleado y Departamento.")
b(sup, "Los KPIs superiores se actualizan según los filtros aplicados: Entradas, Salidas, Tardanzas y Justificadas (sólo las aprobadas).")
img(sup, f"{SUP}/08_reportes_body.jpg", 16.5, "Tabla de marcajes con estado y justificación de cada uno.")

h(sup, "5.2 Exportar CSV para nómina o RRHH", 2)
p(sup, "El botón 'Descargar CSV' genera un archivo con TODOS los registros que cumplan tus filtros (no sólo los primeros 500 visibles). Las columnas incluyen los nuevos campos 'justification_status' y 'rejection_reason'.", size=11)
callout(sup, "SUGERENCIA",
        ["Al procesar el CSV en Excel, aplica un filtro por 'justification_status = rejected' para tener rápidamente la lista de tardanzas que penalizan al empleado.",
         "Combina con la matriz para conciliar antes de cerrar la nómina quincenal o mensual."],
        "EEF5FF", NAVY)
pb(sup)

# ------- 6. Novedades -------
h(sup, "6. Novedades del equipo", 1)
p(sup, "La pantalla 'Novedades' te permite gestionar el ciclo de vacaciones, reposos y permisos de tu equipo.", size=11)
img(sup, f"{SUP}/09_novedades_top.jpg", 16.5, "Novedades — pestañas 'Pendientes / Historial / Todas' y botón 'Nueva novedad'.")

h(sup, "6.1 Pestañas Pendientes · Historial · Todas", 2)
b(sup, "Pendientes: novedades enviadas por tus colaboradores que aún no has decidido. Aquí concentras tu trabajo diario de aprobación.")
b(sup, "Historial: novedades ya decididas (aprobadas o rechazadas). Se conserva la razón y el timestamp de la decisión.")
b(sup, "Todas: vista combinada de pendientes + historial. Útil para búsquedas por empleado o tipo.")

h(sup, "6.2 Aprobar y rechazar en bloque", 2)
n(sup, "Selecciona una o varias novedades marcando la casilla a la izquierda de cada tarjeta.")
n(sup, "Aparecerán los botones 'Aprobar seleccionadas' y 'Rechazar seleccionadas'.")
n(sup, "Al rechazar, ingresa un comentario que verá el colaborador ('Razón del rechazo').")
n(sup, "Al aprobar, opcionalmente puedes añadir un comentario positivo (visible en Historial).")
img(sup, f"{SUP}/10_novedades_all.jpg", 16.5, "Vista 'Todas' de novedades del equipo.")

h(sup, "6.3 Registrar una novedad a nombre de un colaborador", 2)
p(sup, "Cuando un colaborador no puede acceder al sistema (por ejemplo, reposo médico prolongado), tú puedes registrar la novedad por él:", size=11)
n(sup, "Pulsa el botón '+ Nueva novedad' (esquina superior derecha).")
n(sup, "En 'Empleado' selecciona el colaborador correcto.")
n(sup, "Completa Tipo, Fechas y Motivo.")
n(sup, "Envía la solicitud. Quedará automáticamente en 'Pendientes' y podrás aprobarla acto seguido.")
callout(sup, "IMPORTANTE",
        ["Cuando registres una novedad a nombre de otra persona, sé explícito en el motivo (ej. 'Reposo médico entregado en físico a RRHH el 12-feb-2026') para dejar trazabilidad.",
         "Recuerda entregar el original firmado del reposo o permiso a Recursos Humanos."],
        "FFF6E5", AMBER)
pb(sup)

# ------- 7. Buenas prácticas -------
h(sup, "7. Buenas prácticas del supervisor", 1)
b(sup, "Revisa 'Mi Equipo' cada mañana. Las justificaciones pendientes son visibles al abrir la app.")
b(sup, "Decide las justificaciones dentro de 48 horas hábiles. Un colaborador con solicitudes viejas se desmotiva y afecta el clima laboral.")
b(sup, "Registra tus criterios en la razón del rechazo. Es tu evidencia si hay reclamos posteriores.")
b(sup, "Antes de aprobar vacaciones consecutivas, revisa la matriz para verificar cobertura del equipo.")
b(sup, "Usa la Matriz de Asistencia semanalmente para detectar patrones (mismo día tarde, mismo tipo de exceso de descanso) y conversar preventivamente con el colaborador.")
b(sup, "Al cerrar la quincena, exporta el CSV filtrado a 'is_late = true' y valida con Recursos Humanos.")
pb(sup)

# ------- 8. FAQ -------
h(sup, "8. Preguntas frecuentes", 1)
faq = [
    ("¿Puedo modificar una justificación ya decidida?",
     "No. Una vez que has aprobado o rechazado, el status queda fijo en el registro. Si necesitas cambiarlo, contacta al administrador para que ejecute un ajuste manual y deja constancia en el registro por escrito."),
    ("Un colaborador insiste en que su justificación es válida y la rechacé.",
     "Cita al colaborador, escucha sus argumentos y, si aportara evidencia nueva, coordina con Recursos Humanos la anulación de la penalización. La app conserva tu registro original como historial."),
    ("No veo la celda pendiente en Mi Equipo, pero el colaborador dice que ya justificó.",
     "Verifica: (1) que el rango de días incluya la fecha del retraso; (2) que el empleado te tenga como supervisor asignado; (3) refresca con el botón 'Actualizar'."),
    ("¿Cómo veo el equipo sin restricción (vista global)?",
     "Sólo los administradores tienen vista global. Si eres coordinador/gerente/director, sólo verás a quienes tienen tu ID en su campo supervisor_id."),
    ("¿Puedo delegar la aprobación a otro compañero mientras estoy de vacaciones?",
     "Sí. Habla con Recursos Humanos para que reasigne temporalmente supervisor_id de tu equipo. Al terminar tus vacaciones, se revierte."),
    ("¿Se envían notificaciones cuando el colaborador justifica o cuando yo decido?",
     "Actualmente el sistema no envía notificaciones automáticas — es un roadmap próximo. Por ahora, revisa el panel cada mañana."),
]
for q, a in faq:
    p(sup, "P: " + q, bold=True, size=11, color=NAVY)
    p(sup, "R: " + a, size=11)

p(sup, "\nMegaSoft Computación, C.A. — Manual del Supervisor v1.0 — Febrero 2026", size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

sup.save("/app/manual/Manual_Supervisor_MegaSoft_Asistencia.docx")
print("OK Supervisor")


# =============================================================
# MANUAL 2 — ASIGNACIÓN DE HORARIOS
# =============================================================
asg = Document()
setup(asg)
cover(asg,
      "MANUAL DE ASIGNACIÓN DE HORARIOS",
      "Construcción de la matriz para empleados con Horario Especial",
      "Dirigido a: Administradores y supervisores con permiso de planificación")
pb(asg)

h(asg, "Contenido", 1)
for t in [
    "1.  Objetivo del módulo",
    "2.  Conceptos previos: horario fijo vs. Horario Especial",
    "3.  El catálogo de Horarios",
    "     3.1  Bloques, tolerancia general y ventana de justificación",
    "     3.2  Cómo crear un horario nuevo",
    "4.  Asignación de horarios — pantalla principal",
    "     4.1  Elegir rango y empleados",
    "     4.2  Construir la matriz",
    "     4.3  Anatomía de la matriz",
    "5.  Asignar turno a un día",
    "     5.1  Asignación individual (celda por celda)",
    "     5.2  Asignación masiva por selección",
    "     5.3  Asignar novedad (vacaciones / permiso / reposo)",
    "6.  Planificaciones guardadas",
    "     6.1  Guardar una planificación",
    "     6.2  Cargar planificación existente",
    "     6.3  Solapamiento de rangos: qué pasa y cómo resolver",
    "7.  Impacto en la Matriz de Asistencia",
    "     7.1  Cómo aparece un empleado con Horario Especial",
    "     7.2  Cálculo de tardanza según el turno asignado a cada día",
    "     7.3  Faltas, minutos perdidos y novedades",
    "8.  Buenas prácticas y solución de problemas",
]:
    r(asg.add_paragraph(), t, size=11)
pb(asg)

# ------- 1 -------
h(asg, "1. Objetivo del módulo", 1)
p(asg, "El módulo 'Asignación de horarios' permite planificar los turnos rotativos y novedades diarias de los "
       "empleados que NO tienen un horario fijo asignado — llamados en el sistema empleados con 'Horario "
       "Especial'. La matriz que construyes aquí alimenta directamente el 'Reporte matricial' y define la "
       "tolerancia con la que se evaluarán los marcajes de cada día.", size=11)

# ------- 2 -------
h(asg, "2. Conceptos previos: horario fijo vs. Horario Especial", 1)
p(asg, "Cada empleado en el sistema se clasifica en una de dos categorías:", size=11)

t = asg.add_table(rows=3, cols=3)
t.style = "Light Grid Accent 1"
hd = t.rows[0].cells
r(hd[0].paragraphs[0], "Categoría", bold=True, size=10, color=NAVY)
r(hd[1].paragraphs[0], "Cómo se configura", bold=True, size=10, color=NAVY)
r(hd[2].paragraphs[0], "Cómo se evalúa la tardanza", bold=True, size=10, color=NAVY)

r1 = t.rows[1].cells
r(r1[0].paragraphs[0], "Horario Fijo", size=10, bold=True)
r(r1[1].paragraphs[0], "Se le asigna un 'schedule_id' desde la ficha del empleado (Ej.: Día Completo).", size=10)
r(r1[2].paragraphs[0], "Cada día laborable se compara contra los bloques del horario. No requiere planificación.", size=10)

r2 = t.rows[2].cells
r(r2[0].paragraphs[0], "Horario Especial", size=10, bold=True, color=AMBER)
r(r2[1].paragraphs[0], "El empleado NO tiene schedule_id fijo. Aparece como 'elegible' en Asignación de horarios.", size=10)
r(r2[2].paragraphs[0], "Cada día se evalúa según el turno o novedad que le hayas asignado ESE día en la matriz.", size=10)

callout(asg, "CUÁNDO USAR HORARIO ESPECIAL",
        ["Personal de guardia, seguridad o mantenimiento que rota entre turnos día/tarde/noche.",
         "Vendedores externos que alternan visitas a clientes con oficina.",
         "Personal médico o técnico con turnos irregulares (guardias).",
         "Colaboradores con horarios que cambian mes a mes."],
        "EEF5FF", NAVY)
pb(asg)

# ------- 3 -------
h(asg, "3. El catálogo de Horarios", 1)
p(asg, "Antes de asignar horarios en la matriz, deben existir los 'moldes' de horario en el catálogo (menú "
       "lateral 'Horarios'). Cada horario define uno o más bloques de trabajo con su tolerancia.", size=11)
img(asg, f"{ASG}/04_horarios_list.jpg", 16.5, "Catálogo de Horarios: tarjetas con los turnos disponibles.")

h(asg, "3.1 Bloques, tolerancia general y ventana de justificación", 2)
b(asg, "Bloque: rango 'inicio – fin' que representa una jornada. Un horario puede tener 1 bloque (jornada continua) o 2 bloques (jornada partida con descanso al medio).")
b(asg, "Tolerancia general (min): margen para marcar entrada sin ser considerado tarde. Ej.: 10 min → si el bloque inicia a las 08:00 y marcas hasta 08:10 estás 'A tiempo'.")
b(asg, "Ventana de justificación (min): margen adicional a partir del cual el retraso pasa de 'leve' a 'mayor'. Retrasos mayores son OBLIGATORIO justificarlos.")
b(asg, "Sede: opcionalmente, un horario puede estar asociado a una sede específica (útil cuando cada sede tiene una jornada distinta).")

h(asg, "3.2 Cómo crear un horario nuevo", 2)
n(asg, "En el menú lateral, ve a 'Horarios'.")
n(asg, "Pulsa el botón '+ Nuevo horario' en la esquina superior derecha.")
n(asg, "Ingresa un nombre descriptivo (ej. 'Turno Noche · 22:00-06:00').")
n(asg, "Define uno o dos bloques con horas de inicio y fin.")
n(asg, "Establece la 'Tolerancia general' (min) y la 'Justif. (min)' (ventana de justificación).")
n(asg, "Opcionalmente selecciona la sede.")
n(asg, "Presiona 'Guardar'.")
img(asg, f"{ASG}/05_horarios_form.jpg", 13.5, "Formulario para crear un horario en el catálogo.")

callout(asg, "REGLA DE ORO",
        ["Los horarios del catálogo son plantillas: NO se aplican solos a nadie. Deben asignarse (a) como horario fijo en la ficha del empleado, o (b) día a día en la matriz de Asignación de Horarios (para Horario Especial).",
         "Un mismo horario puede reutilizarse en muchos empleados y turnos rotativos."],
        "FFF6E5", AMBER)
pb(asg)

# ------- 4 -------
h(asg, "4. Asignación de horarios — pantalla principal", 1)
p(asg, "Ingresa desde el menú lateral 'Asignación de horarios'. Verás la vista inicial con las opciones para "
       "construir la matriz.", size=11)
img(asg, f"{ASG}/01_asig_top.jpg", 16.5, "Pantalla inicial de Asignación de horarios.")

h(asg, "4.1 Elegir rango y empleados", 2)
b(asg, "Desde / Hasta: rango de días que quieres planificar. Recomendado: quincenal (14 días) o mensual (30 días).")
b(asg, "Empleados sin horario fijo: aparece automáticamente el conteo de personas 'elegibles' (aquellas con Horario Especial). Puedes dejarlo en 'Todos los elegibles' o filtrar a un subconjunto (por departamento, por sede, o eligiéndolos manualmente).")

h(asg, "4.2 Construir la matriz", 2)
n(asg, "Verifica que el rango y los filtros sean correctos.")
n(asg, "Presiona el botón azul 'Construir matriz'.")
n(asg, "La matriz aparecerá vacía por diseño: no hereda datos de planificaciones previas. Debes asignar día por día lo que corresponda.")
img(asg, f"{ASG}/02_asig_matriz.jpg", 16.5, "Matriz construida — cabecera con acciones y filas con empleados.")

h(asg, "4.3 Anatomía de la matriz", 2)
b(asg, "Fila = empleado. Columna = día del rango.")
b(asg, "Celda vacía = sin asignación (será interpretada como 'no obligado a marcar' en la Matriz de Asistencia).")
b(asg, "Celda con turno (ej. 'Día Completo'): el empleado deberá cumplir ese horario ese día.")
b(asg, "Celda con novedad (ej. 'VACACIONES'): día completo cubierto, no se espera marcaje.")
img(asg, f"{ASG}/03_asig_matriz_body.jpg", 16.5, "Detalle del cuerpo de la matriz.")

callout(asg, "REGLA DE LA MATRIZ EN BLANCO",
        ["Al construir una matriz nueva, siempre comienza vacía. Esto evita arrastrar errores de planificaciones anteriores.",
         "Si quieres retomar exactamente una planificación previa, usa la opción 'Cargar planificación existente' (ver sección 6.2)."],
        "EEF5FF", NAVY)
pb(asg)

# ------- 5 -------
h(asg, "5. Asignar turno a un día", 1)

h(asg, "5.1 Asignación individual (celda por celda)", 2)
n(asg, "Haz clic sobre la celda de la fila del empleado y el día deseado.")
n(asg, "Se abrirá un mini-selector con las opciones: turnos disponibles + novedades (Vacaciones, Reposo, Permiso, Remoto, Visita, etc.).")
n(asg, "Elige la opción. La celda se pintará automáticamente con el color o etiqueta correspondiente.")

h(asg, "5.2 Asignación masiva por selección", 2)
n(asg, "Mantén Ctrl (Windows) o ⌘ (Mac) y haz clic en varias celdas para seleccionarlas. También puedes arrastrar el mouse para seleccionar un rango rectangular.")
n(asg, "Verás en la cabecera cuántas celdas están seleccionadas.")
n(asg, "Presiona 'Asignar turno' y elige el turno del catálogo, o 'Asignar novedad' y elige el tipo de novedad. La acción se aplica a TODAS las celdas seleccionadas.")
n(asg, "Usa 'Limpiar' o 'Vaciar' para deshacer selecciones o borrar contenido.")

callout(asg, "TIP DE PRODUCTIVIDAD",
        ["Para planificar un turno rotativo, primero selecciona la columna entera (haz clic en la cabecera del día) y luego 'Asignar turno' → 'Turno Uno'. Repite para los demás turnos.",
         "Al seleccionar la columna, marca la casilla del día en el encabezado; para deseleccionar todo, presiona 'Limpiar'."],
        "EEF5FF", NAVY)

h(asg, "5.3 Asignar novedad (vacaciones / permiso / reposo)", 2)
p(asg, "Una novedad asignada desde la matriz se comporta como una novedad APROBADA de día completo. Es decir:", size=11)
b(asg, "Cubre todo el día: el empleado NO debe marcar.")
b(asg, "No cuenta como falta.")
b(asg, "Se refleja en la Matriz de Asistencia con el color y etiqueta correspondiente (VACACIONES / REPOSO / PERMISO / TRABAJO REMOTO / VISITA).")
b(asg, "Aplica tanto para fechas pasadas como para fechas futuras — te permite planificar vacaciones con anticipación.")

pb(asg)

# ------- 6 -------
h(asg, "6. Planificaciones guardadas", 1)
p(asg, "Una 'planificación' es un snapshot completo de la matriz (rango + empleados + asignaciones) con un nombre que puedes reabrir en cualquier momento.", size=11)

h(asg, "6.1 Guardar una planificación", 2)
n(asg, "Cuando la matriz esté completa, pulsa 'Guardar planificación' en la barra de acciones.")
n(asg, "Dale un nombre descriptivo (ej. 'Guardias Junio 2026 – Sede TBP').")
n(asg, "Confirma. La planificación queda listada en el desplegable 'Cargar planificación existente'.")

h(asg, "6.2 Cargar planificación existente", 2)
n(asg, "Presiona 'Cargar planificación existente' en la esquina superior derecha.")
n(asg, "Selecciona una de la lista. El rango, los empleados y todas las asignaciones se restauran.")
n(asg, "Puedes editarla y volver a guardar (con el mismo nombre para reescribir o con un nombre nuevo para crear una copia).")

h(asg, "6.3 Solapamiento de rangos: qué pasa y cómo resolver", 2)
p(asg, "Si guardas una planificación cuyo rango de fechas se solapa con otra ya guardada, el sistema detectará el conflicto y te ofrecerá dos opciones:", size=11)
b(asg, "Cancelar: se aborta el guardado. Puedes ajustar el rango o cargar la planificación anterior para editarla.")
b(asg, "Reescribir: la(s) planificación(es) anterior(es) se eliminan del catálogo (las asignaciones diarias que ya estaban en la base de datos SÍ se preservan hasta que las sobrescribas explícitamente).")
callout(asg, "IMPORTANTE",
        ["Las asignaciones diarias de la BD son la fuente de verdad. Las 'Planificaciones' son sólo un contenedor para reabrir rápidamente un rango + empleados guardado.",
         "Si eliminas una planificación, las asignaciones diarias que hiciste NO se borran. Sólo se borra el 'contenedor' con nombre.",
         "Para vaciar un día real: entra a la matriz, selecciona la(s) celda(s), y pulsa 'Vaciar'."],
        "FFF6E5", AMBER)
pb(asg)

# ------- 7 -------
h(asg, "7. Impacto en la Matriz de Asistencia", 1)
p(asg, "El resultado de todo tu trabajo en Asignación de horarios se ve reflejado en la 'Matriz de asistencia' (menú lateral).", size=11)
img(asg, f"{ASG}/06_matricial_tipo.jpg", 16.5, "Selector de 'Tipo de horario' — elige 'Horario Especial · turnos rotativos'.")

h(asg, "7.1 Cómo aparece un empleado con Horario Especial", 2)
n(asg, "En 'Reporte matricial de asistencia', selecciona en el filtro 'Tipo de horario' la opción 'Horario Especial · turnos rotativos'.")
n(asg, "Aplica los filtros. Sólo aparecerán los empleados que tengan Horario Especial (los mismos que ves en Asignación de horarios como elegibles).")
n(asg, "Cada día del período mostrará el turno o novedad que asignaste, con los marcajes reales del kiosco superpuestos.")
img(asg, f"{ASG}/07_matricial_especial.jpg", 16.5, "Matriz de asistencia con empleados de Horario Especial y sus novedades aplicadas.")

h(asg, "7.2 Cálculo de tardanza según el turno asignado a cada día", 2)
p(asg, "Éste es el punto CLAVE del módulo: si un empleado Especial marcó a las 08:30 el día que le asignaste 'Turno Uno' (08:00-16:00), se compara contra 08:00 + tolerancia de 'Turno Uno'. Pero si ese mismo empleado el día siguiente tuviera asignado 'Turno Dos' (14:00-22:00), su marca de 14:15 se evaluaría contra 14:00 + tolerancia de 'Turno Dos'.", size=11)
p(asg, "En otras palabras: el sistema recuerda día por día qué horario elegiste y aplica su tolerancia y sus bloques a los marcajes.", size=11, italic=True, color=GREY)

h(asg, "7.3 Faltas, minutos perdidos y novedades", 2)
b(asg, "Si un día se dejó SIN asignar (celda vacía) y NO es día laborable estándar, se marca como 'No laboral' — sin penalización.")
b(asg, "Si un día tiene turno asignado pero el empleado NO marcó, se contabiliza como FALTA en la Matriz.")
b(asg, "Si un día tiene novedad asignada, no se espera marcaje y aparece como VACACIONES/REPOSO/etc. — cuenta al total de 'Novedades', no en 'Faltas'.")
b(asg, "Los minutos de retraso rojos (E1 fuera de tolerancia o E2 > 60 min) suman a 'Minutos perdidos' si NO se justifican o si el supervisor rechaza la justificación.")
pb(asg)

# ------- 8 -------
h(asg, "8. Buenas prácticas y solución de problemas", 1)
b(asg, "Planifica siempre con al menos 1 semana de anticipación. Los empleados agradecen saber su turno.")
b(asg, "Nombra tus planificaciones con formato: '[Grupo] [Mes Año] – [Sede]'. Facilita búsquedas y reutilización.")
b(asg, "Antes de guardar, revisa que ningún colaborador haya quedado con todos los días vacíos por error.")
b(asg, "Si un colaborador se ausenta imprevistamente, asigna 'Reposo' o 'Permiso' el mismo día en la matriz. Esto evita que aparezca como FALTA en el reporte.")
b(asg, "Los cambios en la matriz son inmediatos: no requieren 'publicar'. Sin embargo, la Matriz de Asistencia refleja los cambios recién cuando refrescas los filtros.")
b(asg, "Si un empleado que antes era 'Horario Fijo' pasa a 'Especial', primero limpia su schedule_id en su ficha, luego aparecerá como elegible en Asignación de horarios.")

callout(asg, "SI ALGO NO CUADRA",
        ["1. Verifica en la ficha del empleado que tenga o no schedule_id asignado — según sea fijo o especial.",
         "2. En Asignación de horarios, confirma que aparezca como elegible.",
         "3. En Matriz de asistencia elige 'Horario Especial · turnos rotativos' como Tipo de horario.",
         "4. Si la matriz muestra 'FALTA' en un día que debería ser 'Vacaciones', ve a Asignación de horarios y asigna la novedad correspondiente."],
        "EEF5FF", NAVY)

p(asg, "\nMegaSoft Computación, C.A. — Manual de Asignación de Horarios v1.0 — Febrero 2026",
  size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

asg.save("/app/manual/Manual_Asignacion_Horarios_MegaSoft.docx")
print("OK Asignación")
