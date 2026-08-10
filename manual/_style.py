"""Estilos y utilidades compartidas por los manuales (empleado y supervisor)."""
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def apply_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    sect = doc.sections[0]
    sect.top_margin = Cm(2)
    sect.bottom_margin = Cm(2)
    sect.left_margin = Cm(2.2)
    sect.right_margin = Cm(2.2)


def shade_paragraph(p, hex_color):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def h1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)


def h2(doc, num, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(f"  {num}  ")
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0xFA, 0xCC, 0x15)
    r1.font.size = Pt(14)
    shade_paragraph(p, "0F172A")
    r2 = p.add_run(f"  {text}")
    r2.font.bold = True
    r2.font.size = Pt(16)
    r2.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(4)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(it)


def steps(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(it)


def add_image(doc, path, caption, width_in=6.0):
    doc.add_picture(path, width=Inches(width_in))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def callout(doc, text, kind="tip"):
    palette = {
        "tip":  ("ECFDF5", 0x06, 0x4E, 0x3B),
        "warn": ("FEE2E2", 0x7F, 0x1D, 0x1D),
        "note": ("FEF9C3", 0x71, 0x3F, 0x12),
        "info": ("EFF6FF", 0x1E, 0x3A, 0x8A),
    }
    fill, cr, cg, cb = palette[kind]
    p = doc.add_paragraph()
    shade_paragraph(p, fill)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(cr, cg, cb)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def cover(doc, subtitle_top, title, subtitle_bottom):
    sub = doc.add_paragraph()
    sr = sub.add_run(subtitle_top)
    sr.font.size = Pt(10)
    sr.font.bold = True
    sr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    h1(doc, title)

    kk = doc.add_paragraph()
    kr = kk.add_run(subtitle_bottom)
    kr.font.size = Pt(11)
    kr.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    sep = doc.add_paragraph()
    sep_r = sep.add_run("_" * 90)
    sep_r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
