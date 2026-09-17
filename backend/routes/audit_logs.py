"""Endpoints del Módulo de Pistas de Auditoría (sep-2026).

Consulta con filtros, export XLSX y PDF oficial. RBAC: `auditoria_pistas`.
Regla de inmutabilidad: este módulo solo lee. NO expone endpoints de
UPDATE/DELETE sobre `audit_log`. Los inserts vienen del helper `audit_entity`
llamado desde cada CRUD del sistema.
"""
import base64
import csv
import io
import json
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from deps import (
    api, db, get_current_user, enrich_user_with_permissions,
    HTTPException, Depends, Query, APP_TZ,
)
from fastapi.responses import StreamingResponse

ERR_NO_PERM = "No tienes permiso para consultar las Pistas de Auditoría"

# Módulos del sistema (para dropdown del filtro y validación).
AUDIT_MODULES = [
    "users", "schedules", "schedule_assignments", "assignment_plans",
    "novelties", "visits", "holidays", "access_profiles",
    "sites", "departments", "settings",
]
AUDIT_MODULE_LABELS = {
    "users": "Empleados",
    "schedules": "Horarios",
    "schedule_assignments": "Asignación de turnos",
    "assignment_plans": "Planes de asignación",
    "novelties": "Novedades",
    "visits": "Visitas",
    "holidays": "Días festivos",
    "access_profiles": "Perfiles de acceso",
    "sites": "Sedes",
    "departments": "Departamentos",
    "settings": "Ajustes",
    "auth": "Autenticación",  # Legacy — eventos de login/reset viven aquí
}


async def _require_perm(user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    if user.get("role") == "admin":
        return user
    enriched = await enrich_user_with_permissions(user)
    if (enriched.get("effective_permissions") or {}).get("auditoria_pistas"):
        return user
    raise HTTPException(status_code=403, detail=ERR_NO_PERM)


def _parse_dt(v: Optional[str]) -> Optional[datetime]:
    if not v:
        return None
    try:
        # Acepta ISO con o sin zona; asumimos UTC si viene naïve
        dt = datetime.fromisoformat(v.replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt
    except ValueError:
        raise HTTPException(status_code=400, detail=f"Fecha/hora inválida: {v}")


def _build_query(from_dt: Optional[str], to_dt: Optional[str],
                 modules: Optional[List[str]], actions: Optional[List[str]],
                 user_ids: Optional[List[str]], q: Optional[str]) -> Dict[str, Any]:
    if not from_dt or not to_dt:
        raise HTTPException(status_code=400, detail="El rango de fechas (Desde/Hasta) es obligatorio")
    a = _parse_dt(from_dt)
    b = _parse_dt(to_dt)
    query: Dict[str, Any] = {"timestamp": {"$gte": a, "$lte": b}}
    if modules:
        query["module_name"] = {"$in": modules}
    if actions:
        query["action_type"] = {"$in": [a.upper() for a in actions]}
    if user_ids:
        query["user_id"] = {"$in": user_ids}
    # Nota: el filtro por texto libre (`q`) se aplica en Python — busca
    # dentro del JSON change_detail y en los campos texto (entity_id, email,
    # path, event). No usamos regex Mongo para poder incluir el payload.
    return query


def _matches_q(doc: Dict[str, Any], q: str) -> bool:
    needle = q.lower()
    for k in ("entity_id", "email", "user_id", "path", "event", "action_type", "module_name"):
        v = doc.get(k)
        if v and needle in str(v).lower():
            return True
    if needle in json.dumps(doc.get("change_detail") or {}, default=str, ensure_ascii=False).lower():
        return True
    if needle in json.dumps(doc.get("extra") or {}, default=str, ensure_ascii=False).lower():
        return True
    return False


def _serialize_row(d: Dict[str, Any]) -> Dict[str, Any]:
    d.pop("_id", None)
    ts = d.get("timestamp")
    if isinstance(ts, datetime):
        d["timestamp"] = ts.astimezone(timezone.utc).isoformat()
    return d


async def _search(query: Dict[str, Any], q: Optional[str], limit: int = 5000) -> List[Dict[str, Any]]:
    # Si hay `q`, cargamos hasta 10x del límite para tener margen tras filtrar.
    fetch_limit = min(20000, limit * 10) if q else limit
    docs = await db.audit_log.find(query).sort("timestamp", -1).limit(fetch_limit).to_list(fetch_limit)
    if q:
        docs = [d for d in docs if _matches_q(d, q)][:limit]
    return [_serialize_row(dict(d)) for d in docs]


@api.get("/audit-logs")
async def audit_logs_search(from_dt: Optional[str] = Query(None),
                             to_dt: Optional[str] = Query(None),
                             modules: Optional[List[str]] = Query(None),
                             actions: Optional[List[str]] = Query(None),
                             user_ids: Optional[List[str]] = Query(None),
                             q: Optional[str] = Query(None),
                             limit: int = Query(500, le=5000),
                             user: Dict[str, Any] = Depends(_require_perm)) -> Dict[str, Any]:
    query = _build_query(from_dt, to_dt, modules, actions, user_ids, q)
    rows = await _search(query, q, limit=limit)
    # Sumario para el panel superior de la UI.
    summary: Dict[str, int] = {"CREATE": 0, "UPDATE": 0, "DELETE": 0, "OTHER": 0}
    for r in rows:
        k = (r.get("action_type") or "OTHER").upper()
        summary[k if k in summary else "OTHER"] += 1
    return {
        "from_dt": from_dt, "to_dt": to_dt, "rows_count": len(rows),
        "summary": summary, "rows": rows, "modules": AUDIT_MODULE_LABELS,
    }


@api.get("/audit-logs/modules")
async def audit_logs_modules(user: Dict[str, Any] = Depends(_require_perm)) -> Dict[str, Any]:
    return {"modules": [{"key": k, "label": v} for k, v in AUDIT_MODULE_LABELS.items()]}


@api.get("/audit-logs/export.xlsx")
async def audit_logs_export_xlsx(from_dt: Optional[str] = Query(None),
                                  to_dt: Optional[str] = Query(None),
                                  modules: Optional[List[str]] = Query(None),
                                  actions: Optional[List[str]] = Query(None),
                                  user_ids: Optional[List[str]] = Query(None),
                                  q: Optional[str] = Query(None),
                                  user: Dict[str, Any] = Depends(_require_perm)) -> StreamingResponse:
    query = _build_query(from_dt, to_dt, modules, actions, user_ids, q)
    rows = await _search(query, q, limit=20000)
    settings = await db.settings.find_one({"_id": "company"}) or {}
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoría"
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    dark = PatternFill("solid", fgColor="0F172A")

    ws.merge_cells("A1:H1")
    ws["A1"] = f"{(settings.get('company_name') or 'Mega Soft')} · Pistas de Auditoría"
    ws["A1"].font = Font(bold=True, color="0F172A", size=14)
    ws.merge_cells("A2:H2")
    ws["A2"] = "Gerencia de Seguridad de la Información · Unidad generadora del reporte"
    ws["A2"].font = Font(bold=True, color="B45309", size=10)
    ws.merge_cells("A3:H3")
    ws["A3"] = f"Período: {from_dt} — {to_dt} · {len(rows)} eventos · Generado: {datetime.now(APP_TZ).strftime('%d/%m/%Y %H:%M')}"
    ws["A3"].font = Font(size=9, color="64748B")

    headers = ["Fecha/Hora", "Usuario", "Rol", "IP", "Módulo", "Acción", "Entidad", "Detalle del cambio"]
    widths = [20, 30, 14, 16, 22, 12, 20, 90]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, h in enumerate(headers, start=1):
        c = ws.cell(5, i, h)
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = dark
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    for r_i, r in enumerate(rows, start=6):
        ts = r.get("timestamp") or ""
        if isinstance(ts, str) and "T" in ts:
            try:
                ts = datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(APP_TZ).strftime("%d/%m/%Y %H:%M:%S")
            except Exception:
                pass
        vals = [
            ts, r.get("email") or r.get("user_id") or "—", r.get("role") or "—",
            r.get("ip") or "—",
            AUDIT_MODULE_LABELS.get(r.get("module_name") or "", r.get("module_name") or r.get("event") or "—"),
            r.get("action_type") or (r.get("event") or "—"),
            r.get("entity_id") or "—",
            json.dumps(r.get("change_detail") or r.get("extra") or {}, ensure_ascii=False, default=str)[:32000],
        ]
        for c_i, v in enumerate(vals, start=1):
            cell = ws.cell(r_i, c_i, v)
            cell.font = Font(size=8)
            cell.alignment = Alignment(vertical="top", wrap_text=(c_i == 8))
            cell.border = border

    ws.freeze_panes = "A6"
    buf = io.BytesIO()
    wb.save(buf)
    return StreamingResponse(iter([buf.getvalue()]),
                              media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              headers={"Content-Disposition": f"attachment; filename=auditoria_{from_dt}_a_{to_dt}.xlsx"})


@api.get("/audit-logs/export.pdf")
async def audit_logs_export_pdf(from_dt: Optional[str] = Query(None),
                                 to_dt: Optional[str] = Query(None),
                                 modules: Optional[List[str]] = Query(None),
                                 actions: Optional[List[str]] = Query(None),
                                 user_ids: Optional[List[str]] = Query(None),
                                 q: Optional[str] = Query(None),
                                 user: Dict[str, Any] = Depends(_require_perm)) -> StreamingResponse:
    query = _build_query(from_dt, to_dt, modules, actions, user_ids, q)
    rows = await _search(query, q, limit=5000)
    settings = await db.settings.find_one({"_id": "company"}) or {}
    company = settings.get("company_name") or "Mega Soft"
    logo_b64 = settings.get("logo_base64")

    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=8*mm, rightMargin=8*mm,
                             topMargin=8*mm, bottomMargin=8*mm, title="Pistas de Auditoría")
    styles = getSampleStyleSheet()
    h_title = ParagraphStyle("t", parent=styles["Title"], fontSize=14, textColor=colors.HexColor("#0F172A"))
    h_ger = ParagraphStyle("g", parent=styles["Normal"], fontSize=9, leading=11,
                            textColor=colors.HexColor("#B45309"), fontName="Helvetica-Bold")
    h_sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#64748B"))
    p_cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=6, leading=7.5, textColor=colors.HexColor("#111827"))
    p_head = ParagraphStyle("h", parent=styles["Normal"], fontSize=6.5, leading=8, textColor=colors.white, fontName="Helvetica-Bold")

    story: List[Any] = []
    if logo_b64 and "," in logo_b64:
        try:
            _, b64 = logo_b64.split(",", 1)
            story.append(Image(io.BytesIO(base64.b64decode(b64)), width=30*mm, height=11*mm, hAlign="LEFT"))
        except Exception:
            pass
    story.append(Paragraph("Gerencia de Seguridad de la Información", h_ger))
    story.append(Paragraph("Unidad generadora del reporte", h_sub))
    story.append(Spacer(1, 2*mm))
    story.append(Paragraph(f"{company} · Pistas de Auditoría — Trazabilidad de Cambios", h_title))
    criterios = []
    if modules: criterios.append(f"Módulos: {', '.join(AUDIT_MODULE_LABELS.get(m, m) for m in modules)}")
    if actions: criterios.append(f"Acciones: {', '.join(a.upper() for a in actions)}")
    if user_ids: criterios.append(f"Usuarios: {len(user_ids)}")
    if q: criterios.append(f"Texto: \"{q}\"")
    story.append(Paragraph(f"Período: {from_dt} — {to_dt} · {len(rows)} eventos · "
                            f"Generado: {datetime.now(APP_TZ).strftime('%d/%m/%Y %H:%M')}"
                            + (" · " + " · ".join(criterios) if criterios else ""), h_sub))
    story.append(Spacer(1, 3*mm))

    data = [[Paragraph(h, p_head) for h in ["Fecha/Hora", "Usuario", "IP", "Módulo", "Acción", "Entidad", "Detalle"]]]
    for r in rows:
        ts = r.get("timestamp") or ""
        if isinstance(ts, str) and "T" in ts:
            try:
                ts = datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(APP_TZ).strftime("%d/%m/%Y %H:%M:%S")
            except Exception:
                pass
        detail_txt = json.dumps(r.get("change_detail") or r.get("extra") or {}, ensure_ascii=False, default=str)
        if len(detail_txt) > 500:
            detail_txt = detail_txt[:500] + "…"
        data.append([
            Paragraph(str(ts), p_cell),
            Paragraph(r.get("email") or r.get("user_id") or "—", p_cell),
            Paragraph(r.get("ip") or "—", p_cell),
            Paragraph(AUDIT_MODULE_LABELS.get(r.get("module_name") or "", r.get("module_name") or r.get("event") or "—"), p_cell),
            Paragraph(r.get("action_type") or r.get("event") or "—", p_cell),
            Paragraph(r.get("entity_id") or "—", p_cell),
            Paragraph(detail_txt.replace("<", "&lt;").replace(">", "&gt;"), p_cell),
        ])

    col_widths = [26*mm, 42*mm, 22*mm, 32*mm, 18*mm, 30*mm, 111*mm]
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)
    doc.build(story)
    return StreamingResponse(iter([buf.getvalue()]),
                              media_type="application/pdf",
                              headers={"Content-Disposition": f"attachment; filename=auditoria_{from_dt}_a_{to_dt}.pdf"})
