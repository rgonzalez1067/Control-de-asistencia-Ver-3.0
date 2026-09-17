"""Endpoints de Usuarios (CRUD + importación Excel + selfie + PIN + foto).

Migrado desde server.py durante la Fase A · Iteración 4 (feb-2026).
Cierra la Fase A: `server.py` queda por debajo de las 1000 líneas.
"""
import io
import secrets

from deps import (
    api, db, get_current_user, require_roles,
    now_utc, new_id, strip_mongo_id, sanitize_user, supervisor_scope_ids,
    hash_password, normalize_role,
    LEADER_ROLES,
    UserIn, UserUpdate, SelfieIn, PinIn,
    HTTPException, Depends, UploadFile, File, Request,
    StreamingResponse,
    Any, Dict, List, Optional,
    _load_import_lookups,
    audit_entity,
)

ERR_USER_NOT_FOUND = "Usuario no encontrado"


# ==================================================================
# USERS CRUD (5 endpoints)
# ==================================================================
@api.get("/users")
async def users_list(user: Dict[str, Any] = Depends(get_current_user)) -> List[Dict[str, Any]]:
    q: Dict[str, Any] = {}
    if user.get("role") in LEADER_ROLES:
        q["user_id"] = {"$in": await supervisor_scope_ids(user)}
    docs = await db.users.find(q, {"password_hash": 0, "pin_code_hash": 0,
                                    "selfie_base64": 0, "face_descriptor": 0}).to_list(1000)
    for d in docs:
        strip_mongo_id(d)
    return docs


@api.post("/users")
async def users_create(request: Request, payload: UserIn,
                       actor: Dict[str, Any] = Depends(require_roles("admin"))) -> Dict[str, Any]:
    email = payload.email.lower().strip()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=409, detail="Email ya registrado")
    doc = payload.model_dump(exclude_none=True)
    doc["email"] = email
    doc["user_id"] = new_id("user")
    doc["onboarded"] = False
    doc["created_at"] = now_utc()
    if payload.password:
        from server import validate_password_policy
        validate_password_policy(payload.password)
        doc["password_hash"] = hash_password(payload.password)
        doc["password_updated_at"] = now_utc()
        doc["password_history"] = []
        doc.pop("password", None)
    if payload.pin:
        if not payload.pin.isdigit() or not (4 <= len(payload.pin) <= 8):
            raise HTTPException(status_code=400, detail="PIN debe ser 4-8 dígitos")
        doc["pin_code_hash"] = hash_password(payload.pin)
        doc.pop("pin", None)
    await db.users.insert_one(doc)
    await audit_entity(request, "CREATE", "users", doc["user_id"], after=doc, actor=actor)
    return sanitize_user(doc)


# ------------------------------------------------------------------
# IMPORT (Excel) — DEBE registrarse ANTES de las rutas /users/{user_id}
# ------------------------------------------------------------------
@api.get("/users/import/template")
async def users_import_template(_: Dict[str, Any] = Depends(require_roles("admin"))) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()

    headers = [
        "email", "name", "cedula", "role", "position",
        "department_id", "site_id", "supervisor_id", "schedule_id",
        "password", "kiosk_pin",
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_align = Alignment(horizontal="center", vertical="center")

    ws1 = wb.active
    ws1.title = "Empleados"
    ws1.append(headers)
    for i, _c in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=i)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws1.column_dimensions[cell.column_letter].width = max(14, len(_c) + 4)
    ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet("Ejemplos")
    ws2.append(headers)
    for i, _c in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=i)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws2.column_dimensions[cell.column_letter].width = max(14, len(_c) + 4)
    ws2.freeze_panes = "A2"

    ejemplos = [
        ["jperez@empresa.com", "Juan Pérez", "12345678", "employee",
         "Analista", "Ventas Pyme", "Sede Torre Banco Plaza", "atata@empresa.com", "Día Completo", "Temporal2026*", "1234"],
        ["mrodriguez@empresa.com", "María Rodríguez", "23456789", "Supervisor",
         "Coordinadora", "Recursos Humanos", "Sede Torre Banco Plaza", "", "Día Completo", "", "5678"],
        ["cgomez@empresa.com", "Carlos Gómez", "34567890", "employee",
         "Técnico Soporte", "Soporte y Monitoreo", "Sede Los Chaguaramos", "María Rodríguez", "Turno Uno", "", ""],
    ]
    for row in ejemplos:
        ws2.append(row)

    ws2.append([])
    ws2.append(["NOTAS:"])
    ws2["A" + str(ws2.max_row)].font = Font(bold=True, color="B45309")
    notas = [
        "• Sólo email y name son obligatorios. El resto puede ir vacío.",
    ]
    for n in notas:
        ws2.append([n])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=users_template.xlsx"},
    )


def _get_cell_value(row: tuple, headers: List[str], key: str) -> Optional[str]:
    try:
        idx = headers.index(key)
    except ValueError:
        return None
    if idx >= len(row) or row[idx] is None:
        return None
    s = str(row[idx]).strip()
    return s if s else None


def _resolve_excel_ref(kind: str, value: Optional[str], row_num: int, lookups: tuple, errors: list) -> Optional[str]:
    if not value or not str(value).strip():
        return None
    raw_val = str(value).strip()
    low = raw_val.lower()
    dept_map, site_map, sched_map, sup_by_name, sup_by_email = lookups

    prefix_map = {"department": "dept_", "site": "site_", "schedule": "sch_", "supervisor": "user_"}
    if raw_val.startswith(prefix_map.get(kind, "")):
        return raw_val

    match = None
    if kind == "department":
        match = dept_map.get(low)
    elif kind == "site":
        match = site_map.get(low)
    elif kind == "schedule":
        match = sched_map.get(low)
    elif kind == "supervisor":
        match = sup_by_email.get(low) if "@" in raw_val else sup_by_name.get(low)

    if not match:
        errors.append({"row": row_num, "email": None, "reason": f"{kind.capitalize()} no encontrado: {raw_val!r}"})
    return match


async def _parse_excel_worksheet(file: UploadFile):
    from openpyxl import load_workbook
    raw = await file.read()
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Archivo Excel inválido: {e}") from None
    ws = wb["Empleados"] if "Empleados" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    try:
        headers_row = next(rows)
    except StopIteration:
        raise HTTPException(status_code=400, detail="La hoja está vacía") from None
    headers = [(str(h).strip().lower() if h is not None else "") for h in headers_row]
    if not {"email", "name"}.issubset(set(headers)):
        raise HTTPException(status_code=400, detail="Faltan columnas obligatorias: email, name")
    return headers, rows


@api.post("/users/import/preview")
async def users_import_preview(file: UploadFile = File(...),
                               _: Dict[str, Any] = Depends(require_roles("admin"))) -> Dict[str, Any]:
    headers, rows = await _parse_excel_worksheet(file)
    lookups = await _load_import_lookups()
    to_create, to_update, errors = [], [], []

    for i, row in enumerate(rows, start=2):
        if row is None or all(v is None for v in row):
            continue
        email = (_get_cell_value(row, headers, "email") or "").lower()
        name = _get_cell_value(row, headers, "name")
        if not email or not name:
            errors.append({"row": i, "email": email or None, "reason": "email/name requerido"})
            continue

        candidate = {
            "email": email, "name": name,
            "cedula": _get_cell_value(row, headers, "cedula"),
            "role": normalize_role(_get_cell_value(row, headers, "role")),
            "position": _get_cell_value(row, headers, "position"),
            "department_id": _resolve_excel_ref("department", _get_cell_value(row, headers, "department_id"), i, lookups, errors),
            "site_id": _resolve_excel_ref("site", _get_cell_value(row, headers, "site_id"), i, lookups, errors),
            "supervisor_id": _resolve_excel_ref("supervisor", _get_cell_value(row, headers, "supervisor_id"), i, lookups, errors),
            "schedule_id": _resolve_excel_ref("schedule", _get_cell_value(row, headers, "schedule_id"), i, lookups, errors),
            "kiosk_pin": _get_cell_value(row, headers, "kiosk_pin"),
        }
        existing = await db.users.find_one({"email": email})
        if not existing:
            to_create.append({"row": i, **candidate, "role": candidate["role"] or "employee"})
        else:
            changes = {k: {"from": existing.get(k), "to": v} for k, v in candidate.items() if v is not None and str(existing.get(k) or "") != str(v)}
            to_update.append({"row": i, "email": email, "name": name, "changes": changes, "will_force_name": not changes})

    return {
        "total_rows": len(to_create) + len(to_update) + len(errors),
        "to_create": to_create,
        "to_update": to_update,
        "errors": errors,
    }


@api.post("/users/import")
async def users_import(file: UploadFile = File(...),
                       _: Dict[str, Any] = Depends(require_roles("admin"))) -> Dict[str, Any]:
    headers, rows = await _parse_excel_worksheet(file)
    lookups = await _load_import_lookups()
    created, updated = 0, 0
    created_list, updated_list, errors = [], [], []

    for i, row in enumerate(rows, start=2):
        if row is None or all(v is None for v in row):
            continue
        try:
            email = (_get_cell_value(row, headers, "email") or "").lower()
            name = _get_cell_value(row, headers, "name")
            if not email or not name:
                errors.append({"row": i, "email": email or None, "reason": "email/name requerido"})
                continue

            payload_fields = {
                "cedula": _get_cell_value(row, headers, "cedula"),
                "role": normalize_role(_get_cell_value(row, headers, "role")) if _get_cell_value(row, headers, "role") else None,
                "position": _get_cell_value(row, headers, "position"),
                "department_id": _resolve_excel_ref("department", _get_cell_value(row, headers, "department_id"), i, lookups, errors),
                "site_id": _resolve_excel_ref("site", _get_cell_value(row, headers, "site_id"), i, lookups, errors),
                "supervisor_id": _resolve_excel_ref("supervisor", _get_cell_value(row, headers, "supervisor_id"), i, lookups, errors),
                "schedule_id": _resolve_excel_ref("schedule", _get_cell_value(row, headers, "schedule_id"), i, lookups, errors),
                "kiosk_pin": _get_cell_value(row, headers, "kiosk_pin"),
            }
            set_fields = {k: v for k, v in payload_fields.items() if v is not None}
            set_fields["name"] = name

            existing = await db.users.find_one({"email": email})
            if existing:
                if set_fields:
                    await db.users.update_one({"email": email}, {"$set": set_fields})
                updated += 1
                updated_list.append(email)
            else:
                pw = _get_cell_value(row, headers, "password") or secrets.token_urlsafe(8)
                doc = {
                    "user_id": new_id("user"),
                    "email": email,
                    "name": name,
                    "cedula": payload_fields.get("cedula"),
                    "role": payload_fields.get("role") or "employee",
                    "position": payload_fields.get("position"),
                    "department_id": payload_fields.get("department_id"),
                    "site_id": payload_fields.get("site_id"),
                    "supervisor_id": payload_fields.get("supervisor_id"),
                    "schedule_id": payload_fields.get("schedule_id"),
                    "kiosk_pin": payload_fields.get("kiosk_pin"),
                    "onboarded": False,
                    "created_at": now_utc(),
                    "password_hash": hash_password(pw),
                }
                await db.users.insert_one(doc)
                created += 1
                created_list.append(email)
        except Exception as e:
            errors.append({"row": i, "email": None, "reason": str(e)})

    return {
        "created": created,
        "updated": updated,
        "created_emails": created_list,
        "updated_emails": updated_list,
        "errors": errors,
    }


# ------------------------------------------------------------------
# Rutas con parámetro dinámico {user_id}
# ------------------------------------------------------------------
@api.get("/users/{user_id}")
async def users_get(user_id: str,
                    _: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    u = await db.users.find_one({"user_id": user_id}, {"password_hash": 0, "pin_code_hash": 0})
    if not u:
        raise HTTPException(status_code=404, detail=ERR_USER_NOT_FOUND)
    strip_mongo_id(u)
    return u


@api.put("/users/{user_id}")
async def users_update(request: Request, user_id: str, payload: UserUpdate,
                       actor: Dict[str, Any] = Depends(require_roles("admin"))) -> Dict[str, Any]:
    updates = payload.model_dump(exclude_unset=True)
    if not updates:
        raise HTTPException(status_code=400, detail="Sin cambios")
    before = await db.users.find_one({"user_id": user_id})
    if not before:
        raise HTTPException(status_code=404, detail=ERR_USER_NOT_FOUND)
    await db.users.update_one({"user_id": user_id}, {"$set": updates})
    u = await db.users.find_one({"user_id": user_id}, {"password_hash": 0, "pin_code_hash": 0})
    await audit_entity(request, "UPDATE", "users", user_id, before=before, after=u, actor=actor)
    return strip_mongo_id(u)


@api.post("/users/{user_id}/unlock")
async def users_unlock(request: Request, user_id: str,
                       actor: Dict[str, Any] = Depends(require_roles("admin"))) -> Dict[str, Any]:
    """Admin: desbloqueo manual e inmediato de una cuenta bloqueada por intentos fallidos."""
    before = await db.users.find_one({"user_id": user_id})
    if not before:
        raise HTTPException(status_code=404, detail=ERR_USER_NOT_FOUND)
    await db.users.update_one(
        {"user_id": user_id},
        {"$set": {"failed_login_attempts": 0, "locked_until": None,
                  "unlocked_at": now_utc(), "unlocked_by": actor["user_id"]}},
    )
    after = await db.users.find_one({"user_id": user_id})
    await audit_entity(request, "UPDATE", "users", user_id, before=before, after=after, actor=actor,
                        extra={"action": "unlock_account"})
    return {"ok": True, "user_id": user_id, "unlocked_at": now_utc().isoformat()}


@api.delete("/users/{user_id}")
async def users_delete(request: Request, user_id: str,
                       actor: Dict[str, Any] = Depends(require_roles("admin"))) -> Dict[str, bool]:
    before = await db.users.find_one({"user_id": user_id})
    res = await db.users.delete_one({"user_id": user_id})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail=ERR_USER_NOT_FOUND)
    await audit_entity(request, "DELETE", "users", user_id, before=before, actor=actor)
    return {"ok": True}


@api.post("/users/{user_id}/selfie")
async def users_selfie(user_id: str, payload: SelfieIn,
                       user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, bool]:
    if user["user_id"] != user_id and user["role"] != "admin":
        raise HTTPException(status_code=403, detail="No autorizado")
    updates = {"selfie_base64": payload.selfie_base64, "onboarded": True}
    if payload.face_descriptor is not None:
        updates["face_descriptor"] = payload.face_descriptor
    res = await db.users.update_one({"user_id": user_id}, {"$set": updates})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail=ERR_USER_NOT_FOUND)
    return {"ok": True}


@api.post("/users/{user_id}/pin")
async def users_set_pin(user_id: str, payload: PinIn,
                        user: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, bool]:
    if user["user_id"] != user_id and user["role"] != "admin":
        raise HTTPException(status_code=403, detail="No autorizado")
    if not payload.pin.isdigit() or not (4 <= len(payload.pin) <= 8):
        raise HTTPException(status_code=400, detail="PIN debe ser 4-8 dígitos")
    await db.users.update_one({"user_id": user_id},
                              {"$set": {"pin_code_hash": hash_password(payload.pin)}})
    return {"ok": True}


@api.get("/users/{user_id}/photo")
async def users_get_photo(user_id: str,
                          _: Dict[str, Any] = Depends(get_current_user)) -> Dict[str, Any]:
    u = await db.users.find_one({"user_id": user_id},
                                {"selfie_base64": 1, "onboarded": 1, "name": 1, "_id": 0})
    if not u:
        raise HTTPException(status_code=404, detail=ERR_USER_NOT_FOUND)
    return {
        "user_id": user_id,
        "name": u.get("name"),
        "onboarded": u.get("onboarded", False),
        "selfie_base64": u.get("selfie_base64"),
    }
