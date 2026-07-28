"""Tests for Excel import name-to-ID resolution (department/site/schedule/supervisor)."""
import io
import os
import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
ADMIN_EMAIL = "rgonzalez@megasoft.com.ve"
ADMIN_PASSWORD = "admin123"


@pytest.fixture(scope="module")
def admin_session():
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login",
               json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD}, timeout=20)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text[:300]}"
    data = r.json()
    token = data.get("token") or data.get("access_token")
    if token:
        s.headers.update({"Authorization": f"Bearer {token}"})
    return s


def _build_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Empleados"
    headers = ["email", "name", "cedula", "role", "position",
               "department_id", "site_id", "supervisor_id", "schedule_id",
               "password", "kiosk_pin"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---- 1. Verify existing users have canonical IDs after startup migration ----
def test_existing_users_have_resolved_ids(admin_session):
    r = admin_session.get(f"{BASE_URL}/api/users", timeout=20)
    assert r.status_code == 200
    users = r.json()
    targets = {"narevalo@megasoft.com.ve", "atata@megasoft.com.ve", "asalcedo@megasoft.com.ve"}
    found = {u["email"]: u for u in users if u.get("email") in targets}
    assert targets.issubset(found.keys()), f"missing users: {targets - set(found.keys())}"

    problems = []
    for email, u in found.items():
        site = u.get("site_id")
        sched = u.get("schedule_id")
        # sites and schedules MUST be resolved (per problem statement)
        if site and not site.startswith("site_"):
            problems.append(f"{email}.site_id unresolved: {site!r}")
        if sched and not sched.startswith("sch_"):
            problems.append(f"{email}.schedule_id unresolved: {sched!r}")
        # supervisor: may or may not be resolved (name may not match DB)
        # department: may or may not be resolved (per problem statement)
    assert not problems, "; ".join(problems)


# ---- 2. Preview: valid row resolves + invalid row produces error ----
def test_import_preview_name_resolution(admin_session, tmp_path):
    rows = [
        {"email": "atata@megasoft.com.ve", "name": "Anna Tata Tata",
         "department_id": "Recursos Humanos",
         "site_id": "Sede Torre Banco Plaza",
         "schedule_id": "Día Completo",
         "supervisor_id": "rgonzalez@megasoft.com.ve"},
        {"email": "ghost_test_row@megasoft.com.ve", "name": "Ghost Test",
         "department_id": "Departamento Inventado",
         "site_id": "Sede Torre Banco Plaza",
         "schedule_id": "Día Completo"},
    ]
    buf = _build_xlsx(rows)
    fp = tmp_path / "lookup_test.xlsx"
    fp.write_bytes(buf.getvalue())

    with open(fp, "rb") as f:
        r = admin_session.post(
            f"{BASE_URL}/api/users/import/preview",
            files={"file": ("lookup_test.xlsx", f,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()

    # Invalid row: must be in errors[]
    errs = data.get("errors", [])
    matched = [e for e in errs if "Departamento no encontrado" in e.get("reason", "")
               and "Departamento Inventado" in e.get("reason", "")]
    assert matched, f"expected 'Departamento no encontrado: Departamento Inventado' in errors, got: {errs}"

    # Valid row: should have resolved IDs in the update entry for atata
    upd = [u for u in data.get("to_update", []) if u.get("email") == "atata@megasoft.com.ve"]
    assert upd, f"atata not in to_update: {data.get('to_update')}"
    changes = upd[0].get("changes", {})
    # Verify the 'to' side of each resolved field is a canonical ID
    # (Only present in changes if diff from existing; if already stored, may not appear.
    #  So we check: for each field present, prefix must be correct.)
    def _check_prefix(fld, prefix):
        if fld in changes:
            to_val = changes[fld]["to"]
            assert to_val.startswith(prefix), f"{fld} to={to_val!r} must start with {prefix!r}"
    _check_prefix("department_id", "dept_")
    _check_prefix("site_id", "site_")
    _check_prefix("schedule_id", "sch_")
    _check_prefix("supervisor_id", "user_")


# ---- 3. Import commit: verify DB actually updated with resolved IDs ----
def test_import_commit_updates_resolved_ids(admin_session, tmp_path):
    rows = [
        {"email": "atata@megasoft.com.ve", "name": "Anna Tata Tata",
         "department_id": "Recursos Humanos",
         "site_id": "Sede Torre Banco Plaza",
         "schedule_id": "Día Completo",
         "supervisor_id": "rgonzalez@megasoft.com.ve"},
    ]
    buf = _build_xlsx(rows)
    fp = tmp_path / "commit_test.xlsx"
    fp.write_bytes(buf.getvalue())

    with open(fp, "rb") as f:
        r = admin_session.post(
            f"{BASE_URL}/api/users/import",
            files={"file": ("commit_test.xlsx", f,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=30)
    assert r.status_code == 200, r.text
    result = r.json()
    # errors should not include atata
    for e in result.get("errors", []):
        assert "atata" not in (e.get("email") or "").lower(), f"unexpected err: {e}"

    # Verify via GET /users
    r2 = admin_session.get(f"{BASE_URL}/api/users", timeout=20)
    assert r2.status_code == 200
    atata = next((u for u in r2.json() if u.get("email") == "atata@megasoft.com.ve"), None)
    assert atata, "atata not found after import"
    assert atata.get("department_id", "").startswith("dept_"), atata.get("department_id")
    assert atata.get("site_id", "").startswith("site_"), atata.get("site_id")
    assert atata.get("schedule_id", "").startswith("sch_"), atata.get("schedule_id")
    assert atata.get("supervisor_id", "").startswith("user_"), atata.get("supervisor_id")
    # Expected specific IDs per spec
    assert atata["department_id"] == "dept_30d5899f38", f"got {atata['department_id']}"


# ---- 4. Backward compatibility: raw IDs still accepted ----
def test_import_accepts_raw_ids(admin_session, tmp_path):
    # First fetch current atata values to know a valid dept_id
    r = admin_session.get(f"{BASE_URL}/api/users", timeout=20)
    atata = next((u for u in r.json() if u.get("email") == "atata@megasoft.com.ve"), None)
    assert atata and atata.get("department_id", "").startswith("dept_")
    dept_id = atata["department_id"]

    rows = [
        {"email": "atata@megasoft.com.ve", "name": "Anna Tata Tata",
         "department_id": dept_id},
    ]
    buf = _build_xlsx(rows)
    fp = tmp_path / "raw_id.xlsx"
    fp.write_bytes(buf.getvalue())
    with open(fp, "rb") as f:
        r = admin_session.post(
            f"{BASE_URL}/api/users/import/preview",
            files={"file": ("raw_id.xlsx", f,
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    # No error mentioning department
    dept_errs = [e for e in data.get("errors", []) if "Departamento" in e.get("reason", "")]
    assert not dept_errs, f"raw ID triggered error: {dept_errs}"


# ---- 5. Template Excel: verify Ejemplos sheet has names + notes ----
def test_import_template_uses_names(admin_session):
    r = admin_session.get(f"{BASE_URL}/api/users/import/template", timeout=20)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Ejemplos" in wb.sheetnames, f"sheets: {wb.sheetnames}"
    ws = wb["Ejemplos"]
    all_cells = []
    for row in ws.iter_rows(values_only=True):
        for c in row:
            if c is not None:
                all_cells.append(str(c))
    joined = "\n".join(all_cells)

    for expected in ["Ventas Pyme", "Sede Torre Banco Plaza", "Día Completo",
                     "atata@empresa.com"]:
        assert expected in joined, f"missing example: {expected}"

    # Notes must mention names accepted
    assert "NOMBRE" in joined or "nombre" in joined.lower(), "notes missing name guidance"
