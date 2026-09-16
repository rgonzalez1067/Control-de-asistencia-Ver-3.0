"""Reporte Regulatorio de Visitas — Auditoría de Control de Acceso (sep-2026).

Consolida la trazabilidad completa de cada visita: datos generales, anfitrión,
visitantes (internos/externos) y evidencia biométrica (selfie del Kiosco).

Salidas: JSON (grilla), PDF oficial (reportlab, selfies embebidas) y
XLSX (openpyxl, miniaturas por fila). `include_photos=False` omite la
columna de evidencia fotográfica en todas las salidas.

Reglas de negocio confirmadas con el cliente:
- Sede de la visita = sede del empleado anfitrión (no se persiste en `visits`).
- Hora real de entrada = timestamp de la primera selfie capturada en Kiosco.
- Hora real de salida  = `exit_at` (cierre de la visita).
"""
import base64
from io import BytesIO
from typing import Any, Dict, List, Optional

from PIL import Image as PILImage

from deps import APP_TZ, now_utc, _parse_date_range
from datetime import datetime


# ------------------------------------------------------------------
# Helpers de imagen
# ------------------------------------------------------------------
def _split_b64(data_url: str) -> Optional[str]:
    if not data_url:
        return None
    return data_url.split(",", 1)[1] if "," in data_url else data_url


def _thumb_data_url(data_url: str, max_px: int = 128, quality: int = 60) -> Optional[str]:
    """Devuelve un data URL JPEG reducido para grilla/preview."""
    try:
        b64 = _split_b64(data_url)
        if not b64:
            return None
        img = PILImage.open(BytesIO(base64.b64decode(b64))).convert("RGB")
        img.thumbnail((max_px, max_px))
        out = BytesIO()
        img.save(out, format="JPEG", quality=quality)
        return "data:image/jpeg;base64," + base64.b64encode(out.getvalue()).decode()
    except Exception:
        return None


def _img_buffer(data_url: str, max_px: int) -> Optional[BytesIO]:
    """Buffer JPEG reducido para embeber en PDF/XLSX."""
    try:
        b64 = _split_b64(data_url)
        if not b64:
            return None
        img = PILImage.open(BytesIO(base64.b64decode(b64))).convert("RGB")
        img.thumbnail((max_px, max_px))
        out = BytesIO()
        img.save(out, format="JPEG", quality=75)
        out.seek(0)
        return out
    except Exception:
        return None


def _fmt_dt(dt: Any) -> str:
    if not dt or not isinstance(dt, datetime):
        return "—"
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=__import__("datetime").timezone.utc)
    return dt.astimezone(APP_TZ).strftime("%d/%m/%Y %H:%M")


def _motive_of(v: Dict[str, Any]) -> str:
    if v.get("type") == "laboral":
        return v.get("purpose_label") or v.get("purpose_other") or "—"
    return v.get("purpose_other") or "Visita personal"


# ------------------------------------------------------------------
# Construcción del reporte
# ------------------------------------------------------------------
async def build_visits_report(db, from_date: str, to_date: str,
                              visit_type: Optional[str] = None,
                              site_id: Optional[str] = None,
                              host_user_id: Optional[str] = None,
                              selfies: str = "thumb",
                              include_photos: bool = True) -> Dict[str, Any]:
    rng = _parse_date_range(from_date, to_date)
    q: Dict[str, Any] = {"scheduled_at": rng}
    if visit_type in ("laboral", "personal"):
        q["type"] = visit_type
    if host_user_id:
        q["host_user_id"] = host_user_id

    visits = await db.visits.find(q).sort("scheduled_at", 1).to_list(5000)

    host_ids = list({v.get("host_user_id") for v in visits if v.get("host_user_id")})
    hosts = {u["user_id"]: u async for u in db.users.find({"user_id": {"$in": host_ids}}, {
        "user_id": 1, "name": 1, "cedula": 1, "department_id": 1, "site_id": 1, "_id": 0,
    })}
    depts = {d["department_id"]: d["name"] async for d in db.departments.find({}, {"department_id": 1, "name": 1, "_id": 0})}
    sites = {s["site_id"]: s["name"] async for s in db.sites.find({}, {"site_id": 1, "name": 1, "_id": 0})}

    rows: List[Dict[str, Any]] = []
    visits_included = 0
    for v in visits:
        host = hosts.get(v.get("host_user_id"), {})
        host_site_id = host.get("site_id")
        if site_id and host_site_id != site_id:
            continue
        visits_included += 1
        selfie_list = v.get("selfies") or []
        entry_at = min(
            (s.get("captured_at") for s in selfie_list if s.get("captured_at")),
            default=None,
        )
        base = {
            "visit_id": v.get("visit_id"),
            "type": v.get("type"),
            "status": v.get("status"),
            "scheduled_at": v.get("scheduled_at"),
            "entry_at": entry_at,
            "exit_at": v.get("exit_at"),
            "duration_minutes": v.get("duration_minutes"),
            "site_id": host_site_id,
            "site_name": sites.get(host_site_id, "—"),
            "host_name": v.get("host_name") or host.get("name") or "—",
            "host_cedula": host.get("cedula") or "—",
            "host_department": depts.get(host.get("department_id") or "", "—"),
            "company_name": v.get("company_name"),
            "motive": _motive_of(v),
            "observations": v.get("observations"),
        }
        for idx, vis in enumerate(v.get("visitors") or []):
            s = next((x for x in selfie_list if x.get("visitor_index") == idx), None)
            selfie_b64 = (s or {}).get("selfie_base64")
            row = {
                **base,
                "visitor_name": vis.get("name") or "—",
                "visitor_cedula": vis.get("cedula") or ("Menor de edad" if vis.get("is_minor") else "—"),
                "visitor_phone": vis.get("phone"),
                "visitor_kind": "Interno" if vis.get("kind") == "internal" else "Externo",
                "is_minor": bool(vis.get("is_minor")),
                "has_selfie": bool(selfie_b64),
            }
            if include_photos and selfies == "thumb" and selfie_b64:
                row["selfie_thumb"] = _thumb_data_url(selfie_b64)
            elif include_photos and selfies == "full" and selfie_b64:
                row["selfie"] = selfie_b64
            rows.append(row)

    return {
        "from_date": from_date,
        "to_date": to_date,
        "generated_at": now_utc().isoformat(),
        "include_photos": include_photos,
        "visits_count": visits_included,
        "visitors_count": len(rows),
        "rows": rows,
    }


# ------------------------------------------------------------------
# Exportación XLSX
# ------------------------------------------------------------------
def export_visits_xlsx(report: Dict[str, Any], include_photos: bool = True) -> bytes:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Visitas Realizadas"

    headers = ["Código", "Fecha/Hora Programada", "Entrada Real", "Salida Real",
               "Sede", "Tipo", "Anfitrión", "Cédula Anfitrión", "Departamento",
               "Visitante", "Cédula Visitante", "Clasificación", "Empresa / Motivo"]
    widths = [16, 17, 17, 17, 20, 10, 28, 15, 20, 28, 15, 12, 30]
    if include_photos:
        headers.append("Foto")
        widths.append(13)

    ws.append(headers)
    head_fill = PatternFill("solid", fgColor="0F172A")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for r_i, row in enumerate(report["rows"], start=2):
        values = [
            row["visit_id"], _fmt_dt(row["scheduled_at"]), _fmt_dt(row["entry_at"]),
            _fmt_dt(row["exit_at"]), row["site_name"],
            "Laboral" if row["type"] == "laboral" else "Personal",
            row["host_name"], row["host_cedula"], row["host_department"],
            row["visitor_name"], row["visitor_cedula"], row["visitor_kind"],
            (row["company_name"] or "") + (f" · {row['motive']}" if row.get("motive") else "") or row.get("motive") or "—",
        ]
        if include_photos:
            values.append("")
        ws.append(values)
        if include_photos:
            ws.row_dimensions[r_i].height = 62
        for c in range(1, len(headers) + 1):
            ws.cell(row=r_i, column=c).alignment = Alignment(vertical="center", wrap_text=True)
            ws.cell(row=r_i, column=c).font = Font(size=8)
        if include_photos and row.get("selfie"):
            buf = _img_buffer(row["selfie"], max_px=96)
            if buf:
                img = XLImage(buf)
                img.anchor = f"{get_column_letter(len(headers))}{r_i}"
                ws.add_image(img)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------
# Exportación PDF (formato oficial de auditoría)
# ------------------------------------------------------------------
def export_visits_pdf(report: Dict[str, Any], company_name: str = "Mega Soft",
                      logo_base64: Optional[str] = None, include_photos: bool = True) -> bytes:
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=8 * mm, rightMargin=8 * mm,
                            topMargin=8 * mm, bottomMargin=8 * mm,
                            title="Reporte de Visitas Realizadas")

    styles = getSampleStyleSheet()
    h_title = ParagraphStyle("t", parent=styles["Title"], fontSize=14, textColor=colors.HexColor("#0F172A"))
    h_sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=8, textColor=colors.HexColor("#64748B"))
    h_ger = ParagraphStyle("g", parent=styles["Normal"], fontSize=9, leading=11,
                           textColor=colors.HexColor("#B45309"), fontName="Helvetica-Bold",
                           spaceBefore=2)
    p_cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=6.5, leading=8, textColor=colors.HexColor("#111827"))
    p_head = ParagraphStyle("h", parent=styles["Normal"], fontSize=6.5, leading=8,
                            textColor=colors.white, fontName="Helvetica-Bold")

    story: List[Any] = []
    if logo_base64 and "," in logo_base64:
        try:
            _, b64 = logo_base64.split(",", 1)
            story.append(Image(BytesIO(base64.b64decode(b64)), width=30 * mm, height=11 * mm, hAlign="LEFT"))
        except Exception:
            pass
    story.append(Paragraph("Gerencia de Seguridad de la Información", h_ger))
    story.append(Paragraph("Unidad generadora del reporte", ParagraphStyle(
        "g2", parent=styles["Normal"], fontSize=6.5, textColor=colors.HexColor("#94A3B8"))))
    story.append(Spacer(1, 2 * mm))
    story.append(Paragraph(f"{company_name} · Reporte de Visitas Realizadas — Auditoría de Control de Acceso", h_title))
    story.append(Paragraph(
        f"Período: {report['from_date']} — {report['to_date']} · "
        f"{report['visits_count']} visitas · {report['visitors_count']} visitantes · "
        f"Evidencia fotográfica: {'incluida' if include_photos else 'no incluida'} · "
        f"Generado: {datetime.now(APP_TZ).strftime('%d/%m/%Y %H:%M')}", h_sub))
    story.append(Spacer(1, 3 * mm))

    headers = ["Código", "Programada", "Entrada Real", "Salida Real", "Sede", "Tipo",
               "Anfitrión", "Departamento", "Visitante", "Clasif.", "Empresa / Motivo"]
    col_widths = [16 * mm, 21 * mm, 21 * mm, 21 * mm, 22 * mm, 13 * mm,
                  34 * mm, 22 * mm, 34 * mm, 13 * mm, 32 * mm]
    if include_photos:
        headers.append("Foto")
        col_widths.append(19 * mm)

    data = [[Paragraph(h, p_head) for h in headers]]

    for row in report["rows"]:
        anf = f"{row['host_name']}<br/><font color='#64748B'>{row['host_cedula']}</font>"
        vis = f"{row['visitor_name']}<br/><font color='#64748B'>{row['visitor_cedula']}</font>"
        emp_mot = row["company_name"] or ""
        if row.get("motive"):
            emp_mot = f"{emp_mot} · {row['motive']}" if emp_mot else row["motive"]
        line = [
            Paragraph(str(row["visit_id"] or "").replace("visit_", ""), p_cell),
            Paragraph(_fmt_dt(row["scheduled_at"]), p_cell),
            Paragraph(_fmt_dt(row["entry_at"]), p_cell),
            Paragraph(_fmt_dt(row["exit_at"]), p_cell),
            Paragraph(row["site_name"] or "—", p_cell),
            Paragraph("Laboral" if row["type"] == "laboral" else "Personal", p_cell),
            Paragraph(anf, p_cell),
            Paragraph(row["host_department"] or "—", p_cell),
            Paragraph(vis, p_cell),
            Paragraph(row["visitor_kind"], p_cell),
            Paragraph(emp_mot or "—", p_cell),
        ]
        if include_photos:
            img_cell: Any = ""
            if row.get("selfie"):
                ibuf = _img_buffer(row["selfie"], max_px=300)
                if ibuf:
                    img_cell = Image(ibuf, width=17 * mm, height=17 * mm)
            line.append(img_cell)
        data.append(line)

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)
    doc.build(story)
    return buf.getvalue()
